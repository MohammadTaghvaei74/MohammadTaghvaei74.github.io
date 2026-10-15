from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent
payload = json.loads((ROOT / "data/history.json").read_text(encoding="utf-8"))
out = ROOT / "data/market-score-history.xlsx"

wb = Workbook()
readme = wb.active
readme.title = "راهنما"
readme.sheet_view.rightToLeft = True
readme_rows = [
    ["آرشیو امتیاز نبض بازار"],
    ["روزهای کامل و تأییدشده", payload["valid_day_count"]],
    ["تعداد نمادها", len(payload["symbols"])],
    ["تعداد ردیف‌های امتیاز", payload["row_count"]],
    ["تعریف mean", "میانگین امتیازهای ثبت‌شده نماد در همان روز"],
    ["تعریف median", "میانه امتیازهای ثبت‌شده نماد در همان روز"],
    ["کیفیت review", "مقدار از دامنه عادی نمایش نمودار خارج است؛ حذف نشده و نیازمند بررسی است."],
    ["منبع", "scores_history.db / v2_days + v2_daily_scores"],
]
for row in readme_rows:
    readme.append(row)
readme.column_dimensions["A"].width = 28
readme.column_dimensions["B"].width = 95
readme["A1"].font = Font(size=16, bold=True, color="C9F27C")
readme["A1"].fill = PatternFill("solid", fgColor="0B2028")
readme.merge_cells("A1:B1")

summary = wb.create_sheet("روزهای معتبر")
summary.sheet_view.rightToLeft = True
summary.append(["تاریخ", "ارزش معاملات (تومان)", "وضعیت", "منبع تأیید"])
for day in payload["valid_days"]:
    summary.append([day["day"], day["turnover"], "کامل و تأییدشده", day["source"]])
summary.column_dimensions["A"].width = 16
summary.column_dimensions["B"].width = 23
summary.column_dimensions["C"].width = 20
summary.column_dimensions["D"].width = 105

scores = wb.create_sheet("امتیاز روزانه")
scores.sheet_view.rightToLeft = True
headers = ["تاریخ", "نماد", "تعداد نمونه", "میانگین", "میانه", "کمینه", "بیشینه", "آخرین", "کیفیت"]
scores.append(headers)
for row in payload["rows"]:
    scores.append([row.get("day"), row.get("symbol"), row.get("count"), row.get("mean"), row.get("median"), row.get("min"), row.get("max"), row.get("last"), "نیازمند بررسی" if row.get("quality") == "review" else "معتبر"])
for col, width in {"A": 16, "B": 18, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 18}.items():
    scores.column_dimensions[col].width = width
scores.freeze_panes = "A2"

for ws in [summary, scores]:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="14333A")
        cell.alignment = Alignment(horizontal="right")
    ws.auto_filter.ref = ws.dimensions

scores_table = Table(displayName="DailyScores", ref=f"A1:I{scores.max_row}")
scores_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
scores.add_table(scores_table)
summary_table = Table(displayName="ValidDays", ref=f"A1:D{summary.max_row}")
summary_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
summary.add_table(summary_table)

for row in scores.iter_rows(min_row=2, min_col=4, max_col=8):
    for cell in row:
        cell.number_format = "0.00"
for cell in summary["B"][1:]:
    cell.number_format = "#,##0"

wb.properties.title = "Market Pulse Score History"
wb.properties.subject = "Verified daily stock score archive"
wb.save(out)
print(out)
print(out.stat().st_size)
print(scores.max_row - 1)
print(summary.max_row - 1)
